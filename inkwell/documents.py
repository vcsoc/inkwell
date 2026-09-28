"""Local Documents editor. File access is confined to the user's Documents directory."""

import base64
import ctypes
import errno
import html
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import textwrap
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

import nh3
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import Response
from lxml import html as lxml_html
from markdown_it import MarkdownIt
from PIL import Image
from pydantic import BaseModel, Field

from . import store

router = APIRouter(prefix="/api/documents")
TEXT_EXT = {".txt", ".md", ".mdx", ".csv", ".tsv", ".yaml", ".yml", ".json", ".xml", ".toml", ".ini", ".log", ".py", ".js", ".css"}
RICH_EXT = {".docx", ".doc", ".odt", ".rtf", ".html", ".htm"}
FORMATS = TEXT_EXT | RICH_EXT | {".pdf"}
MAX_FILE = 40_000_000
MAX_HTML = 8_000_000
MAX_IMAGE = 5_000_000
HTML_TAGS = {"p", "div", "br", "strong", "b", "em", "i", "u", "s", "strike", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre", "code", "ul", "ol", "li", "table", "thead", "tbody", "tr", "td", "th", "a", "img", "font", "span", "hr", "sub", "sup"}
HTML_ATTRS = {"a": {"href", "title"}, "img": {"src", "alt", "width", "height"}, "font": {"color", "size", "face"}, "td": {"colspan", "rowspan"}, "th": {"colspan", "rowspan"}, "p": {"align"}, "div": {"align"}, "h1": {"align"}, "h2": {"align"}, "h3": {"align"}}
IMAGE_URL = re.compile(r"^data:image/(png|jpeg|webp|gif);base64,([A-Za-z0-9+/=]+)$", re.I)
PDF_PASSWORDS = {}
PDF_PASSWORD_LOCK = threading.Lock()
PDF_PASSWORD_TTL = 30 * 60


def root():
    override = os.getenv("INKWELL_DOCUMENTS_DIR")
    if override:
        folder = Path(override).expanduser()
    else:
        folder = Path.home() / "Documents"
        config = Path.home() / ".config" / "user-dirs.dirs"
        try:
            match = re.search(r'^XDG_DOCUMENTS_DIR="([^"\n]+)"', config.read_text(), re.M)
            if match:
                folder = Path(match.group(1).replace("$HOME", str(Path.home()), 1)).expanduser()
        except OSError:
            pass
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def resolve(path="", *, expect=None):
    if not isinstance(path, str) or len(path) > 1600 or "\\" in path or "\x00" in path:
        raise HTTPException(422, "Invalid document path")
    parts = PurePosixPath(path).parts
    if path.startswith("/") or any(part in ("..", ".") or part.startswith(".") for part in parts):
        raise HTTPException(403, "Document path is outside Documents")
    base = root()
    target = base.joinpath(*parts)
    if not target.resolve().is_relative_to(base):
        raise HTTPException(403, "Document path is outside Documents")
    current = base
    for part in parts:
        current /= part
        if current.is_symlink():
            raise HTTPException(403, "Links are not available in Documents")
    if expect and (not target.exists() or (target.is_file() if expect == "dir" else target.is_dir())):
        raise HTTPException(404, "Document or folder not found")
    return target


def relative(path):
    return path.relative_to(root()).as_posix() if path != root() else ""


def name_ok(name):
    if not isinstance(name, str) or name in ("", ".", "..") or len(name) > 160 or any(c in name for c in "/\\\x00") or name.startswith(".") or any(ord(c) < 32 for c in name):
        raise HTTPException(422, "Choose a valid document or folder name")
    return name


def check_file(path):
    file = resolve(path, expect="file")
    if file.suffix.lower() not in FORMATS:
        raise HTTPException(415, "This file type cannot be edited")
    if file.stat().st_size > MAX_FILE:
        raise HTTPException(413, "Document is larger than 40 MB")
    return file


def revision(file):
    stat = file.stat()
    return f"{stat.st_mtime_ns}:{stat.st_size}"


def unlock_pdf(document, file, password=None):
    """Keep a successful password in process memory only, keyed to the current file revision."""
    if not document.needs_pass:
        return
    key = (relative(file), revision(file))
    if password is None:
        with PDF_PASSWORD_LOCK:
            cached = PDF_PASSWORDS.get(key)
            password = cached[0] if cached and cached[1] > time.monotonic() else None
            if cached and password is None:
                PDF_PASSWORDS.pop(key, None)
    if password is None:
        raise HTTPException(423, "Document password required")
    if not document.authenticate(password):
        raise HTTPException(401, "Incorrect document password")
    with PDF_PASSWORD_LOCK:
        PDF_PASSWORDS[key] = (password, time.monotonic() + PDF_PASSWORD_TTL)
        if len(PDF_PASSWORDS) > 100:
            for stale in list(PDF_PASSWORDS)[:50]:
                PDF_PASSWORDS.pop(stale, None)


def recent():
    try:
        paths = json.loads(store.setting("recent_documents", "[]"))
    except ValueError:
        return []
    valid = []
    for path in paths[:30]:
        try:
            file = check_file(path)
            valid.append({"path": path, "name": file.name, "kind": file.suffix.lower()[1:]})
        except (HTTPException, OSError):
            continue
    return valid[:12]


def touch_recent(path):
    paths = [item["path"] for item in recent() if item["path"] != path]
    store.set_setting("recent_documents", json.dumps([path, *paths][:12]))


def sanitize(content):
    if len(content.encode("utf-8")) > MAX_HTML:
        raise HTTPException(413, "Document content exceeds 8 MB")
    # Reject remote images, SVG, unsafe links and CSS URLs, including unquoted attributes.
    def filter_attribute(tag, attr, value):
        if tag == "img" and attr == "src" and (not IMAGE_URL.fullmatch(value) or len(value) > MAX_IMAGE * 2):
            return None
        if tag == "a" and attr == "href" and not re.match(r"^(https://|mailto:)", value, re.I):
            return None
        if attr in ("width", "height") and (not value.isdigit() or int(value) > 2000):
            return None
        if tag == "font" and attr == "color" and not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
            return None
        if tag == "font" and attr == "size" and value not in {str(i) for i in range(1, 8)}:
            return None
        if tag == "font" and attr == "face" and value not in {"Arial", "Georgia", "Times New Roman", "Courier New"}:
            return None
        return value

    return nh3.clean(content, tags=HTML_TAGS, attributes=HTML_ATTRS, attribute_filter=filter_attribute, allowed_classes={"img": {"doc-float-left", "doc-float-right"}}, url_schemes={"https", "mailto", "data"}, clean_content_tags={"script", "style", "iframe", "object", "svg", "form", "template"})


def image_bytes(value):
    match = IMAGE_URL.fullmatch(value or "")
    if not match:
        raise HTTPException(422, "Choose a PNG, JPEG, WebP or GIF image")
    try:
        data = base64.b64decode(match.group(2), validate=True)
        if len(data) > MAX_IMAGE:
            raise ValueError("Image too large")
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        return data
    except (ValueError, OSError):
        raise HTTPException(422, "Invalid image") from None


def inline_html(paragraph, part):
    fragments = []
    for element in paragraph._p:
        if element.tag == qn("w:hyperlink"):
            relationship = part.rels.get(element.get(qn("r:id")))
            label = html.escape("".join(node.text or "" for node in element.iter(qn("w:t"))))
            if relationship and re.match(r"^(https://|mailto:)", relationship.target_ref, re.I):
                target = html.escape(relationship.target_ref, quote=True)
                fragments.append(f'<a href="{target}">{label}</a>')
            else:
                fragments.append(label)
            continue
        if element.tag != qn("w:r"):
            continue
        run = Run(element, paragraph)
        text = html.escape(run.text).replace("\n", "<br>")
        if run.bold:
            text = "<strong>" + text + "</strong>"
        if run.italic:
            text = "<em>" + text + "</em>"
        if run.underline:
            text = "<u>" + text + "</u>"
        if run.font.strike:
            text = "<s>" + text + "</s>"
        if run.font.color.rgb is not None or run.font.name or run.font.size:
            attributes = []
            if run.font.color.rgb is not None:
                attributes.append(f'color="#{run.font.color.rgb}"')
            if run.font.name:
                attributes.append(f'face="{html.escape(run.font.name, quote=True)}"')
            if run.font.size:
                size = min(range(1, 8), key=lambda option: abs((8, 10, 12, 14, 18, 24, 32)[option - 1] - run.font.size.pt))
                attributes.append(f'size="{size}"')
            text = "<font " + " ".join(attributes) + ">" + text + "</font>"
        fragments.append(text)
        for blip in run.element.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            if rid not in part.related_parts:
                continue
            data = part.related_parts[rid].blob
            if len(data) > MAX_IMAGE:
                continue
            try:
                with Image.open(io.BytesIO(data)) as image:
                    media = Image.MIME.get(image.format, "")
            except OSError:
                continue
            if media in ("image/png", "image/jpeg", "image/gif", "image/webp"):
                extent = run.element.xpath(".//wp:extent")
                try:
                    width = min(720, max(40, round(int(extent[0].get("cx")) * 96 / 914400))) if extent else 480
                except (TypeError, ValueError):
                    width = 480
                fragments.append(f'<img src="data:{media};base64,{base64.b64encode(data).decode()}" alt="Embedded image" width="{width}">')
    return "".join(fragments)


def docx_to_html(file):
    document = Document(str(file))
    output = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            para = Paragraph(element, document)
            style = para.style.name if para.style else ""
            heading = re.fullmatch(r"Heading ([1-6])", style)
            tag = f"h{heading.group(1)}" if heading else "p"
            output.append(f"<{tag}>{inline_html(para, document.part)}</{tag}>")
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            output.append("<table><tbody>")
            for row in table.rows:
                output.append("<tr>")
                for cell in row.cells:
                    output.append("<td>" + "".join("<p>" + inline_html(p, document.part) + "</p>" for p in cell.paragraphs) + "</td>")
                output.append("</tr>")
            output.append("</tbody></table>")
    return sanitize("".join(output))


def inline_to_docx(parent, node, bold=False, italic=False, underline=False, strike=False, font_color=None, font_face=None, font_size=None):
    tag = node.tag.lower() if isinstance(node.tag, str) else ""
    bold = bold or tag in ("b", "strong")
    italic = italic or tag in ("i", "em")
    underline = underline or tag == "u"
    strike = strike or tag in ("s", "strike")
    if tag == "font":
        font_color = node.get("color") or font_color
        font_face = node.get("face") or font_face
        font_size = node.get("size") or font_size

    def add_text(text):
        run = parent.add_run(text)
        run.bold, run.italic, run.underline = bold, italic, underline
        run.font.strike = strike
        if font_color:
            run.font.color.rgb = RGBColor.from_string(font_color[1:])
        if font_face:
            run.font.name = font_face
        if font_size:
            run.font.size = Pt((8, 10, 12, 14, 18, 24, 32)[int(font_size) - 1])

    if tag == "a" and re.match(r"^(https://|mailto:)", node.get("href", ""), re.I):
        relationship = parent.part.relate_to(node.get("href"), RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = "".join(node.itertext())
        run.append(text)
        hyperlink.append(run)
        parent._p.append(hyperlink)
        return
    if tag == "img":
        data = image_bytes(node.get("src"))
        try:
            width = max(40, min(int(node.get("width", "480")), 1400))
        except ValueError:
            width = 480
        parent.add_run().add_picture(io.BytesIO(data), width=int(min(width / 96, 7.5) * 914400))
        return
    if tag == "br":
        parent.add_run().add_break()
    if node.text:
        add_text(node.text)
    for child in node:
        inline_to_docx(parent, child, bold, italic, underline, strike, font_color, font_face, font_size)
        if child.tail:
            add_text(child.tail)


def html_to_docx(content):
    parsed = lxml_html.fragment_fromstring(sanitize(content), create_parent="div")
    document = Document()
    if parsed.text and parsed.text.strip():
        document.add_paragraph(parsed.text)
    for item in parsed:
        tag = item.tag.lower() if isinstance(item.tag, str) else ""
        if tag == "table":
            rows = item.xpath("./tbody/tr|./tr|./thead/tr")
            if not rows:
                continue
            columns = max(len(row.xpath("./td|./th")) for row in rows)
            if not columns:
                continue
            table = document.add_table(rows=len(rows), cols=min(columns, 30))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row.xpath("./td|./th")[:30]):
                    inline_to_docx(table.cell(i, j).paragraphs[0], cell)
        elif tag in ("ul", "ol"):
            for child in item.xpath("./li"):
                paragraph = document.add_paragraph(style="List Bullet" if tag == "ul" else "List Number")
                inline_to_docx(paragraph, child)
        else:
            heading = re.fullmatch(r"h([1-6])", tag)
            paragraph = document.add_paragraph(style=f"Heading {heading.group(1)}" if heading else None)
            if tag == "blockquote":
                paragraph.style = "Quote"
            if item.get("align") == "center":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif item.get("align") == "right":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            inline_to_docx(paragraph, item)
        if item.tail and item.tail.strip():
            document.add_paragraph(item.tail)
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


def office_convert(source, suffix):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise HTTPException(503, "Install LibreOffice to edit legacy .doc, .odt or .rtf documents")
    with tempfile.TemporaryDirectory(prefix="inkwell-document-") as directory:
        profile = Path(directory) / "profile"
        output = Path(directory) / "output"
        output.mkdir()
        try:
            result = subprocess.run([executable, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", suffix.lstrip("."), "--outdir", str(output), str(source)], capture_output=True, timeout=35, cwd=directory)
            converted = output / (source.stem + suffix)
            if result.returncode or not converted.is_file() or converted.stat().st_size > MAX_FILE:
                raise HTTPException(422, "LibreOffice could not convert this document")
            return converted.read_bytes()
        except subprocess.TimeoutExpired:
            raise HTTPException(504, "Document conversion timed out") from None


def atomic_write(path, data, *, exclusive=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".inkwell-", delete=False) as output:
        temp = Path(output.name)
        try:
            os.chmod(temp, 0o600)
            output.write(data)
        except Exception:
            temp.unlink(missing_ok=True)
            raise
    try:
        if exclusive:
            # Linking a complete temp file publishes it without replacing a user file,
            # including one created between choosing the name and writing the output.
            try:
                os.link(temp, path)
            except FileExistsError:
                raise HTTPException(409, "An output file already has that name; retry") from None
        else:
            temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def output_path(source, label):
    base = source.with_name(source.stem + "." + label + source.suffix)
    candidate = base
    number = 2
    while candidate.exists() or candidate.is_symlink():
        candidate = source.with_name(f"{source.stem}.{label}-{number}{source.suffix}")
        number += 1
    return candidate


@router.get("/tree")
def tree(path: str = ""):
    folder = resolve(path, expect="dir")
    children = []
    try:
        for child in folder.iterdir():
            if child.name.startswith(".") or child.is_symlink() or not (child.is_dir() or child.suffix.lower() in FORMATS):
                continue
            stat = child.stat()
            children.append({"name": child.name, "path": relative(child), "directory": child.is_dir(), "size": stat.st_size if child.is_file() else None, "modified": stat.st_mtime_ns})
            if len(children) >= 2000:
                break
    except PermissionError:
        raise HTTPException(403, "Cannot read this folder") from None
    children.sort(key=lambda item: (not item["directory"], item["name"].casefold()))
    return {"name": folder.name, "path": relative(folder), "children": children}


@router.get("/search")
def search_documents(q: str = Query(min_length=1, max_length=120)):
    """Search only visible names beneath Documents; never follow hidden folders or symlinks."""
    matches = []
    scanned = 0
    needle = q.casefold()
    for folder, directories, files in os.walk(root(), followlinks=False):
        directories[:] = [name for name in directories if not name.startswith(".") and not (Path(folder) / name).is_symlink()]
        for name in [*directories, *files]:
            scanned += 1
            if scanned > 30_000 or len(matches) >= 200:
                break
            path = Path(folder) / name
            if name.startswith(".") or path.is_symlink() or (path.is_file() and path.suffix.lower() not in FORMATS) or needle not in name.casefold():
                continue
            matches.append({"name": name, "path": relative(path), "directory": path.is_dir(), "modified": path.stat().st_mtime_ns})
        if scanned > 30_000 or len(matches) >= 200:
            break
    return sorted(matches, key=lambda item: (not item["directory"], item["name"].casefold()))


class MoveItem(BaseModel):
    path: str
    destination: str = ""


@router.post("/move")
def move_document(data: MoveItem):
    source = resolve(data.path)
    folder = resolve(data.destination, expect="dir")
    if source == root() or not source.exists() or source.is_symlink() or source.name.startswith("."):
        raise HTTPException(404, "Document or folder not found")
    if source.is_file() and source.suffix.lower() not in FORMATS:
        raise HTTPException(415, "This file type cannot be moved here")
    if source.is_dir() and folder.is_relative_to(source):
        raise HTTPException(422, "A folder cannot be moved into itself or a descendant")
    target = folder / source.name
    if target == source:
        return {"path": relative(source), "unchanged": True}
    if target.exists() or target.is_symlink():
        raise HTTPException(409, "A file or folder already has that name")
    old_path = relative(source)
    previous = [item["path"] for item in recent()]
    try:
        source.rename(target)
    except OSError as error:
        raise HTTPException(422, "Could not move this item between these folders") from error
    updated = [relative(target) + path[len(old_path):] if path == old_path or path.startswith(old_path + "/") else path for path in previous]
    store.set_setting("recent_documents", json.dumps(updated[:12]))
    return {"path": relative(target), "old_path": old_path}


class CopyItems(BaseModel):
    paths: list[str] = Field(min_length=1, max_length=128)
    destination: str = ""


@router.post("/copy")
def copy_documents(data: CopyItems):
    """Copy selected Documents items without following links or replacing existing files."""
    folder = resolve(data.destination, expect="dir")
    sources = [resolve(path) for path in data.paths]
    if len(set(sources)) != len(sources):
        raise HTTPException(422, "Select each item only once")
    for source in sources:
        if source == root() or not source.exists() or source.is_symlink():
            raise HTTPException(404, "Document or folder not found")
        if source.is_file() and source.suffix.lower() not in FORMATS:
            raise HTTPException(415, "This file type cannot be copied here")
        if source.is_dir() and folder.is_relative_to(source):
            raise HTTPException(422, "A folder cannot be copied into itself or a descendant")
        if any(source != other and source.is_relative_to(other) for other in sources):
            raise HTTPException(422, "Select either a folder or its contents, not both")
        if source.is_dir():
            for base, directories, files in os.walk(source, followlinks=False):
                if any((Path(base) / name).is_symlink() for name in [*directories, *files]):
                    raise HTTPException(403, "Folders containing links cannot be copied")
    copied = []
    for source in sources:
        # Temporary hidden staging keeps interrupted copies out of the visible tree.
        with tempfile.TemporaryDirectory(prefix=".inkwell-copy-", dir=folder) as staging:
            staged = Path(staging) / source.name
            if source.is_dir():
                shutil.copytree(source, staged, symlinks=True)
            else:
                shutil.copy2(source, staged, follow_symlinks=False)
            stem = source.name if source.is_dir() else source.stem
            suffix = "" if source.is_dir() else source.suffix
            number = 1
            while True:
                label = " (copy)" if number == 1 else f" (copy {number})"
                target = folder / (stem + label + suffix)
                if not target.exists() and not target.is_symlink():
                    break
                number += 1
            try:
                if staged.is_dir():
                    # Linux renameat2(NOREPLACE) publishes a staged tree atomically,
                    # even if another process creates the target after our check.
                    rename = ctypes.CDLL(None, use_errno=True).renameat2
                    if rename(-100, os.fsencode(staged), -100, os.fsencode(target), 1):
                        code = ctypes.get_errno()
                        raise OSError(code, os.strerror(code))
                else:
                    os.link(staged, target, follow_symlinks=False)
            except OSError as error:
                if error.errno == errno.EEXIST:
                    raise HTTPException(409, "A copied item already has that name; retry") from error
                raise
            copied.append(relative(target))
    return {"paths": copied}


@router.get("/recent")
def list_recent():
    return recent()


class UnlockDocument(BaseModel):
    path: str
    password: str = Field(max_length=1024)


@router.post("/unlock")
def unlock_document(data: UnlockDocument):
    return read_document(data.path, data.password)


@router.get("/open")
def open_document(path: str):
    return read_document(path)


def read_document(path, password=None):
    file = check_file(path)
    kind = file.suffix.lower()
    result = {"path": relative(file), "name": file.name, "revision": revision(file), "kind": kind[1:]}
    if kind == ".pdf":
        try:
            with pymupdf.open(file) as document:
                unlock_pdf(document, file, password)
                result.update(mode="pdf", pages=[{"width": round(p.rect.width, 2), "height": round(p.rect.height, 2)} for p in document])
        except (pymupdf.FileDataError, pymupdf.EmptyFileError):
            raise HTTPException(422, "Invalid PDF") from None
    elif kind in RICH_EXT:
        if kind in (".html", ".htm"):
            result.update(mode="rich", content=sanitize(file.read_text(encoding="utf-8-sig")))
        else:
            if kind == ".docx":
                source = file
                result.update(mode="rich", content=docx_to_html(source))
            else:
                with tempfile.TemporaryDirectory(prefix="inkwell-document-") as directory:
                    converted = Path(directory) / (file.stem + ".docx")
                    converted.write_bytes(office_convert(file, ".docx"))
                    result.update(mode="rich", content=docx_to_html(converted))
    else:
        try:
            result.update(mode="text", content=file.read_text(encoding="utf-8-sig"))
        except UnicodeError:
            raise HTTPException(422, "This text file is not UTF-8") from None
    touch_recent(relative(file))
    return result


@router.get("/page")
def pdf_page(path: str, page: int = 0, width: int = Query(900, ge=100, le=1800)):
    file = check_file(path)
    if file.suffix.lower() != ".pdf":
        raise HTTPException(415, "Page previews require a PDF")
    try:
        with pymupdf.open(file) as document:
            unlock_pdf(document, file)
            if page < 0 or page >= len(document):
                raise HTTPException(422, "Invalid PDF page")
            target = document[page]
            if target.rect.width <= 0 or target.rect.height <= 0 or width * width * target.rect.height / target.rect.width > 8_000_000:
                raise HTTPException(413, "PDF page is too large to preview safely")
            pixmap = target.get_pixmap(matrix=pymupdf.Matrix(width / target.rect.width, width / target.rect.width), alpha=False)
            return Response(pixmap.tobytes("png"), media_type="image/png", headers={"Cache-Control": "no-store"})
    except (pymupdf.FileDataError, pymupdf.EmptyFileError):
        raise HTTPException(422, "Invalid PDF") from None


class NewItem(BaseModel):
    path: str = ""
    name: str = Field(max_length=160)


@router.post("/folder")
def create_folder(data: NewItem):
    folder = resolve(data.path, expect="dir") / name_ok(data.name)
    if folder.exists():
        raise HTTPException(409, "A file or folder already has that name")
    folder.mkdir(mode=0o700)
    return {"path": relative(folder)}


@router.post("/file")
def create_file(data: NewItem):
    folder = resolve(data.path, expect="dir")
    name = name_ok(data.name)
    file = folder / name
    if file.suffix.lower() not in FORMATS:
        raise HTTPException(415, "Choose a supported document format")
    if file.exists():
        raise HTTPException(409, "A file already has that name")
    suffix = file.suffix.lower()
    if suffix == ".pdf":
        pdf = pymupdf.open()
        pdf.new_page()
        content = pdf.tobytes()
        pdf.close()
    elif suffix == ".docx":
        content = html_to_docx("<p></p>")
    elif suffix in (".doc", ".odt", ".rtf"):
        with tempfile.TemporaryDirectory(prefix="inkwell-document-") as directory:
            source = Path(directory) / "new.docx"
            source.write_bytes(html_to_docx("<p></p>"))
            content = office_convert(source, suffix)
    else:
        content = b""
    # Exclusive create: never overwrite a file created in between the check and write.
    try:
        with file.open("xb") as output:
            os.chmod(file, 0o600)
            output.write(content)
    except FileExistsError:
        raise HTTPException(409, "A file already has that name") from None
    return {"path": relative(file)}


@router.post("/import")
async def import_file(request: Request, path: str = "", name: str = ""):
    folder = resolve(path, expect="dir")
    name = name_ok(name)
    if Path(name).suffix.lower() not in FORMATS:
        raise HTTPException(415, "Choose a supported document format")
    if request.headers.get("content-length") and int(request.headers["content-length"]) > MAX_FILE:
        raise HTTPException(413, "Document exceeds 40 MB")
    incoming = bytearray()
    async for chunk in request.stream():
        if len(incoming) + len(chunk) > MAX_FILE:
            raise HTTPException(413, "Document exceeds 40 MB")
        incoming.extend(chunk)
    if not incoming:
        raise HTTPException(422, "Choose a non-empty document")
    data = bytes(incoming)
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise HTTPException(422, "Invalid PDF file")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise HTTPException(422, "Invalid DOCX file")
    if suffix in TEXT_EXT | {".html", ".htm"}:
        try:
            data.decode("utf-8-sig")
        except UnicodeError:
            raise HTTPException(422, "Text documents must be UTF-8") from None
    file = folder / name
    try:
        with file.open("xb") as output:
            os.chmod(file, 0o600)
            output.write(data)
    except FileExistsError:
        raise HTTPException(409, "A file already has that name") from None
    return {"path": relative(file)}


class SaveContent(BaseModel):
    path: str
    revision: str
    content: str = Field(max_length=MAX_HTML)


@router.put("/content")
def save_content(data: SaveContent):
    file = check_file(data.path)
    if revision(file) != data.revision:
        raise HTTPException(409, "Document changed on disk. Reopen it before saving")
    kind = file.suffix.lower()
    if kind == ".pdf":
        raise HTTPException(415, "Use PDF tools to edit and export a new PDF")
    if kind in TEXT_EXT:
        content = data.content.encode("utf-8")
    elif kind in (".html", ".htm"):
        content = sanitize(data.content).encode("utf-8")
    else:
        source_data = html_to_docx(data.content)
        if kind == ".docx":
            content = source_data
        else:
            with tempfile.TemporaryDirectory(prefix="inkwell-document-") as directory:
                source = Path(directory) / (file.stem + ".docx")
                source.write_bytes(source_data)
                content = office_convert(source, kind)
    if len(content) > MAX_FILE:
        raise HTTPException(413, "Document exceeds 40 MB")
    atomic_write(file, content)
    touch_recent(relative(file))
    return {"path": relative(file), "revision": revision(file)}


class PdfOperation(BaseModel):
    path: str
    action: str
    page: int = 0
    rect: list[float] | None = None
    image: str | None = None
    text: str | None = Field(default=None, max_length=5000)
    font_size: int = Field(default=12, ge=6, le=32)


def pdf_rect(page, rect):
    if not rect or len(rect) != 4 or not all(isinstance(v, (int, float)) and v == v and abs(v) < 100_000 for v in rect):
        raise HTTPException(422, "Select a valid area on the page")
    value = pymupdf.Rect(rect)
    if value.is_empty or value.is_infinite or not page.rect.contains(value) or value.width < 2 or value.height < 2:
        raise HTTPException(422, "Selected area must be inside the page")
    return value


class PdfTextArea(BaseModel):
    path: str
    page: int = 0
    rect: list[float]


@router.post("/pdf-text")
def pdf_text_in_area(data: PdfTextArea):
    file = check_file(data.path)
    if file.suffix.lower() != ".pdf":
        raise HTTPException(415, "PDF text selection requires a PDF document")
    try:
        with pymupdf.open(file) as document:
            unlock_pdf(document, file)
            if data.page < 0 or data.page >= len(document):
                raise HTTPException(422, "Page does not exist")
            page = document[data.page]
            return {"text": page.get_textbox(pdf_rect(page, data.rect))[:5000]}
    except (pymupdf.FileDataError, pymupdf.EmptyFileError):
        raise HTTPException(422, "Invalid PDF") from None


@router.post("/pdf")
def edit_pdf(data: PdfOperation):
    file = check_file(data.path)
    if file.suffix.lower() != ".pdf":
        raise HTTPException(415, "PDF tools require a PDF document")
    if data.action not in {"redact", "image", "signature", "text", "replace_text", "compress"}:
        raise HTTPException(422, "Unknown PDF action")
    try:
        with pymupdf.open(file) as document:
            unlock_pdf(document, file)
            if data.action == "compress":
                # Re-encode oversized images, remove unused objects and compress streams.
                document.rewrite_images(dpi_threshold=180, dpi_target=144, quality=68)
                document.subset_fonts()
            else:
                if data.page < 0 or data.page >= len(document):
                    raise HTTPException(422, "Page does not exist")
                page = document[data.page]
                rect = pdf_rect(page, data.rect)
                if data.action == "redact":
                    page.add_redact_annot(rect, fill=(0, 0, 0))
                    page.apply_redactions(images=2, graphics=1, text=0)
                elif data.action in ("image", "signature"):
                    image = image_bytes(data.image)
                    with Image.open(io.BytesIO(image)) as source:
                        image_width, image_height = source.size
                    # PyMuPDF's keep_proportion can still paint a stretched image into
                    # a non-square rect. Fit the actual rectangle before insertion.
                    scale = min(rect.width / image_width, rect.height / image_height)
                    width, height = image_width * scale, image_height * scale
                    fitted = pymupdf.Rect(
                        rect.x0 + (rect.width - width) / 2,
                        rect.y0 + (rect.height - height) / 2,
                        rect.x0 + (rect.width + width) / 2,
                        rect.y0 + (rect.height + height) / 2,
                    )
                    page.insert_image(fitted, stream=image, keep_proportion=False)
                elif data.action in ("text", "replace_text"):
                    if not data.text or not data.text.strip():
                        raise HTTPException(422, "Enter text to place")
                    if data.action == "replace_text":
                        # Permanently remove covered text in the NEW copy before inserting the replacement.
                        page.add_redact_annot(rect, fill=(1, 1, 1))
                        page.apply_redactions(images=0, graphics=0, text=0)
                    if page.insert_textbox(rect, data.text, fontsize=data.font_size, fontname="helv") < 0:
                        raise HTTPException(422, "The selected area is too small for this text")
            output = document.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    except (pymupdf.FileDataError, pymupdf.EmptyFileError):
        raise HTTPException(422, "Invalid PDF") from None
    if data.action == "compress" and len(output) >= file.stat().st_size:
        return {"path": relative(file), "saved_bytes": 0, "message": "The PDF is already as small as this compression can make it"}
    result = output_path(file, {"signature": "signed-image", "image": "image", "redact": "redacted", "text": "annotated", "replace_text": "replaced-text", "compress": "compressed"}[data.action])
    atomic_write(result, output, exclusive=True)
    touch_recent(relative(result))
    return {"path": relative(result), "saved_bytes": file.stat().st_size - len(output) if data.action == "compress" else 0}


class DigitalSignature(BaseModel):
    path: str
    certificate: str = Field(max_length=3_000_000)
    password: str = Field(default="", max_length=1024)
    page: int = 0
    rect: list[float] | None = None


@router.post("/sign")
def sign_pdf(data: DigitalSignature):
    """Sign a new PDF copy; PKCS#12 certificate and passphrase are never stored."""
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    from pyhanko.sign import fields, signers

    file = check_file(data.path)
    if file.suffix.lower() != ".pdf":
        raise HTTPException(415, "Digital signatures require a PDF")
    try:
        certificate = base64.b64decode(data.certificate, validate=True)
        if not certificate or len(certificate) > 2_000_000:
            raise ValueError("Invalid certificate")
        signer = signers.SimpleSigner.load_pkcs12_data(certificate, other_certs=(), passphrase=data.password.encode())
        if signer is None:
            raise ValueError("Cannot read certificate")
        with pymupdf.open(file) as document:
            unlock_pdf(document, file)
            if document.needs_pass:
                raise HTTPException(422, "Digital signing encrypted PDFs is not yet supported")
            if data.page < 0 or data.page >= len(document):
                raise ValueError("Invalid PDF page")
            page = document[data.page]
            selected = pdf_rect(page, data.rect) if data.rect else pymupdf.Rect(page.rect.width - 220, page.rect.height - 95, page.rect.width - 20, page.rect.height - 25)
            box = (selected.x0, page.rect.height - selected.y1, selected.x1, page.rect.height - selected.y0)
        field = "InkwellSignature" + datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S%f")
        writer = IncrementalPdfFileWriter(io.BytesIO(file.read_bytes()))
        destination = io.BytesIO()
        signers.sign_pdf(writer, signature_meta=signers.PdfSignatureMetadata(field_name=field), signer=signer, new_field_spec=fields.SigFieldSpec(sig_field_name=field, on_page=data.page, box=box), output=destination)
        result = output_path(file, "digitally-signed")
        atomic_write(result, destination.getvalue(), exclusive=True)
        touch_recent(relative(result))
        return {"path": relative(result)}
    except HTTPException:
        raise
    except (ValueError, OSError, KeyError, TypeError) as error:
        raise HTTPException(422, "Could not sign the PDF. Check its PKCS#12 certificate and password") from error


class MarkdownPreview(BaseModel):
    content: str = Field(max_length=MAX_HTML)


@router.post("/preview")
def preview_markdown(data: MarkdownPreview):
    rendered = MarkdownIt("commonmark", {"html": False}).render(data.content)
    return {"html": sanitize(rendered)}


def pdf_content(file):
    """Produce a printable PDF, retaining office layout when LibreOffice is available."""
    kind = file.suffix.lower()
    if kind == ".pdf":
        with pymupdf.open(file) as document:
            unlock_pdf(document, file)
            return document.tobytes(encryption=pymupdf.PDF_ENCRYPT_NONE) if document.needs_pass else file.read_bytes()
    if kind in RICH_EXT and (shutil.which("libreoffice") or shutil.which("soffice")):
        converted = office_convert(file, ".pdf")
        if len(converted) > MAX_FILE:
            raise HTTPException(413, "PDF export exceeds 40 MB")
        return converted
    if kind == ".docx":
        document = Document(str(file))
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] + [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
    elif kind in (".doc", ".odt", ".rtf"):
        raise HTTPException(503, "Install LibreOffice to export this legacy document as PDF")
    else:
        try:
            text = file.read_text(encoding="utf-8-sig")
            if kind in (".html", ".htm"):
                cleaned = sanitize(text)
                text = lxml_html.fromstring(cleaned).text_content() if cleaned.strip() else ""
        except (UnicodeError, ValueError):
            raise HTTPException(422, "Text documents must be UTF-8") from None
    if kind in (".md", ".mdx"):
        rendered = MarkdownIt("commonmark", {"html": False}).render(text)
        text = lxml_html.fromstring(rendered).text_content() if rendered.strip() else ""
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    y = 52
    for original in text.splitlines():
        for line in textwrap.wrap(original.expandtabs(4), width=92, replace_whitespace=False, drop_whitespace=False) or [""]:
            if y > 800:
                page = pdf.new_page(width=595, height=842)
                y = 52
            page.insert_text((42, y), line, fontsize=10, fontname="helv")
            y += 15
    result = pdf.tobytes(garbage=3, deflate=True)
    pdf.close()
    if len(result) > MAX_FILE:
        raise HTTPException(413, "PDF export exceeds 40 MB")
    return result


class DocumentPath(BaseModel):
    path: str


@router.post("/export-pdf")
def export_pdf(data: DocumentPath):
    source = check_file(data.path)
    if source.suffix.lower() == ".pdf":
        return {"path": relative(source), "message": "This document is already a PDF"}
    output = pdf_content(source)
    target = source.with_name(source.stem + ".export.pdf")
    index = 2
    while target.exists() or target.is_symlink():
        target = source.with_name(f"{source.stem}.export-{index}.pdf")
        index += 1
    atomic_write(target, output, exclusive=True)
    touch_recent(relative(target))
    return {"path": relative(target)}


@router.get("/print-preview")
def print_preview(path: str):
    output = pdf_content(check_file(path))
    with pymupdf.open(stream=output, filetype="pdf") as document:
        if not len(document):
            raise HTTPException(422, "Document has no printable pages")
        page = document[0]
        width = 500
        if page.rect.width <= 0 or width * width * page.rect.height / page.rect.width > 5_000_000:
            raise HTTPException(413, "Page is too large to preview")
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(width / page.rect.width, width / page.rect.width), alpha=False)
        return Response(pixmap.tobytes("png"), media_type="image/png", headers={"Cache-Control": "no-store"})


def available_printers():
    if not shutil.which("lpstat"):
        return [], None
    try:
        status = subprocess.run(["lpstat", "-p", "-d"], capture_output=True, text=True, timeout=4)
    except (OSError, subprocess.TimeoutExpired):
        return [], None
    names = re.findall(r"^printer ([A-Za-z0-9_.-]{1,128}) ", status.stdout, re.M)
    default = re.search(r"^system default destination: ([A-Za-z0-9_.-]{1,128})$", status.stdout, re.M)
    return names, default.group(1) if default else None


@router.get("/printers")
def list_printers():
    names, default = available_printers()
    return {"printers": names, "default": default}


class PrintOptions(BaseModel):
    path: str
    printer: str
    copies: int = Field(default=1, ge=1, le=99)
    pages: str = Field(default="", max_length=120)
    paper: str = "A4"
    orientation: str = "portrait"
    color: str = "color"
    duplex: str = "none"
    scaling: str = "fit"
    margins: str = "default"


@router.post("/print")
def print_document(data: PrintOptions):
    if not shutil.which("lp"):
        raise HTTPException(503, "CUPS printing is not available on this system")
    printers, _ = available_printers()
    if data.printer not in printers:
        raise HTTPException(422, "Choose an available printer")
    if data.paper not in {"A4", "Letter", "Legal"} or data.orientation not in {"portrait", "landscape"} or data.color not in {"color", "monochrome"} or data.duplex not in {"none", "long", "short"} or data.scaling not in {"fit", "actual"} or data.margins not in {"default", "narrow", "none"}:
        raise HTTPException(422, "Choose supported print settings")
    if data.pages and not re.fullmatch(r"[1-9][0-9]*(?:-[1-9][0-9]*)?(?:,[1-9][0-9]*(?:-[1-9][0-9]*)?)*", data.pages):
        raise HTTPException(422, "Use page numbers like 1-3,5")
    output = pdf_content(check_file(data.path))
    args = ["lp", "-d", data.printer, "-n", str(data.copies), "-o", f"media={data.paper}", "-o", f"orientation-requested={4 if data.orientation == 'landscape' else 3}", "-o", f"print-color-mode={data.color}", "-o", f"sides={'two-sided-long-edge' if data.duplex == 'long' else 'two-sided-short-edge' if data.duplex == 'short' else 'one-sided'}"]
    if data.pages:
        args.extend(["-o", "page-ranges=" + data.pages])
    if data.scaling == "fit":
        args.extend(["-o", "fit-to-page"])
    if data.margins != "default":
        margin = 0 if data.margins == "none" else 18
        for side in ("left", "right", "top", "bottom"):
            args.extend(["-o", f"page-{side}={margin}"])
    with tempfile.NamedTemporaryFile(dir=root(), prefix=".inkwell-print-", suffix=".pdf") as temporary:
        temporary.write(output)
        temporary.flush()
        try:
            result = subprocess.run([*args, temporary.name], capture_output=True, text=True, timeout=25)
        except (OSError, subprocess.TimeoutExpired):
            raise HTTPException(503, "Could not contact the printer") from None
    if result.returncode:
        raise HTTPException(502, "Printer rejected this job; check its status and settings")
    return {"message": result.stdout.strip()[:240] or "Print job submitted"}


@router.get("/download")
def download(path: str):
    file = check_file(path)
    from fastapi.responses import FileResponse

    return FileResponse(file, filename=file.name, media_type="application/octet-stream", headers={"X-Content-Type-Options": "nosniff"})
